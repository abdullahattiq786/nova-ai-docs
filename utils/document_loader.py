"""Extract textual content and useful source metadata without persistent uploads."""
from io import BytesIO
import csv
import io
import os
import tempfile

from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as WordDocument
from openpyxl import load_workbook
from pptx import Presentation


def load_document(filename: str, data: bytes) -> list[Document]:
    ext = os.path.splitext(filename)[1].lower()
    docs = []

    def add(text: str, **metadata):
        if text and text.strip():
            docs.append(Document(page_content=text.strip(), metadata={"source": filename, **metadata}))

    if ext == ".pdf":
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")
        for page_number, page in enumerate(reader.pages, 1):
            add(page.extract_text() or "", page=page_number)
    elif ext == ".docx":
        doc = WordDocument(BytesIO(data))
        for paragraph in doc.paragraphs:
            add(paragraph.text, section="paragraph")
        for table_num, table in enumerate(doc.tables, 1):
            for row_num, row in enumerate(table.rows, 1):
                add(" | ".join(cell.text for cell in row.cells), table=table_num, row=row_num)
    elif ext in (".txt", ".md"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Text must be UTF-8 encoded.")
        add(text)
    elif ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV must be UTF-8 encoded.")
        for row_num, row in enumerate(csv.reader(io.StringIO(text)), 1):
            add(" | ".join(row), row=row_num)
    elif ext == ".xlsx":
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in wb:
                for row_num, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    add(" | ".join("" if cell is None else str(cell) for cell in row), sheet=sheet.title, row=row_num)
        finally:
            wb.close()
    elif ext == ".pptx":
        presentation = Presentation(BytesIO(data))
        for slide_num, slide in enumerate(presentation.slides, 1):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
            add("\n".join(parts), slide=slide_num)
    else:
        raise ValueError("Unsupported document type.")

    if not docs:
        raise ValueError("No extractable text found. Scanned images and image-only PDFs need OCR.")
    return docs
